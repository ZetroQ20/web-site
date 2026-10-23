#!/usr/bin/env python3
"""Dump slide text and tables from a .pptx or .pdf as Markdown, one section per slide/page."""
import sys
from pathlib import Path


def md_table(rows):
    rows = [[(c or "").replace("\n", " ").replace("|", "\\|").strip() for c in r] for r in rows if r]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def from_pptx(path):
    from pptx import Presentation

    def shapes(group):
        for shape in group:
            if shape.shape_type == 6:  # group shape
                yield from shapes(shape.shapes)
            else:
                yield shape

    for i, slide in enumerate(Presentation(path).slides, 1):
        print(f"\n## Slide {i}\n")
        for shape in shapes(slide.shapes):
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in p.runs).strip()
                    if text:
                        print("  " * p.level + "- " + text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                print("\n" + md_table(rows) + "\n")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                print(f"\n> Speaker notes: {notes}")


def from_pdf(path):
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            print(f"\n## Page {i}\n")
            text = (page.extract_text() or "").strip()
            print(text if text else "[no text layer - render this page and read it as an image]")
            for table in page.extract_tables():
                print("\n" + md_table(table) + "\n")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: extract_slides.py <file.pptx|file.pdf>")
    src = Path(sys.argv[1])
    {".pptx": from_pptx, ".pdf": from_pdf}.get(src.suffix.lower(), lambda _: sys.exit(f"unsupported: {src.suffix}"))(src)
