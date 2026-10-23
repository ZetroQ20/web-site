#!/usr/bin/env python3
"""Build a .docx (and optionally .pdf) from bilingual Markdown notes with pandoc, then tidy tables and spacing."""
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


def shade(cell, fill="D9E2F3"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def polish(docx_path):
    doc = Document(docx_path)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(54)
    for style in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if style in [s.name for s in doc.styles]:
            s = doc.styles[style]
            s.font.size = Pt(10.5)
            s.paragraph_format.space_after = Pt(3)
    for table in doc.tables:
        set_borders(table)
        for cell in table.rows[0].cells:
            shade(cell)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
    doc.save(docx_path)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 1:
        sys.exit("usage: build_notes.py <notes.md> [--pdf]")
    md = Path(args[0])
    out = md.with_suffix(".docx")
    subprocess.run(["pandoc", str(md), "-f", "markdown+pipe_tables", f"--resource-path={md.parent}", "-o", str(out)],
                   check=True)
    polish(out)
    print(out)
    if "--pdf" in sys.argv:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out.parent), str(out)],
                       check=True, stdout=subprocess.DEVNULL)
        print(out.with_suffix(".pdf"))


if __name__ == "__main__":
    main()
